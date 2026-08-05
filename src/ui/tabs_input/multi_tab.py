""" 다중 케이스 실행 설정 모달리스 창 (MultiCaseRunWindow) """

import os
import tempfile
import shutil
import json
import openpyxl

from PySide6.QtWidgets import (
    QDialog, QWidget, QVBoxLayout, QLabel, QPushButton, QComboBox,
    QTableWidget, QTableWidgetItem, QAbstractItemView, QFileDialog,
    QMessageBox,
)
from PySide6.QtGui import QBrush, QColor
from PySide6.QtCore import Qt
from PySide6.QtCore import QCoreApplication

class MultiCaseRunWindow(QDialog):
    """ 모달리스 다중 케이스 실행 설정 창 """

    parameter_map = {
        "WaveHs"        : { "Target_File": ".sea", "File_Symbol": "Hs", "Description": "Hs" },
        "WaveTp"        : { "Target_File": ".sea", "File_Symbol": "Tp", "Description": "Tp" },
        "PropagationDir": { "Target_File": ".inf", "File_Symbol": "WD", "Description": "Wind Direction" },
        "WaveDir"       : { "Target_File": ".sea", "File_Symbol": "AD", "Description": "Wave Direction" },
        "CurrSSDir"     : { "Target_File": ".sea", "File_Symbol": "CD", "Description": "Current Direction" },
        "NacYaw"        : { "Target_File": ".ela", "File_Symbol": "YE", "Description": "Yaw Error" },
        "FileName_BTS"  : { "Target_File": ".inf", "File_Symbol": "WS", "Description": "Wind Seed" },
        "WaveSeed(1)"   : { "Target_File": ".sea", "File_Symbol": "AS", "Description": "Wave Seed" },
    }

    def __init__(self, parent, file_path, ref_name, excel_path):
        super().__init__(parent, Qt.Window)
        self.file_path = file_path
        self.ref_name = ref_name
        self.excel_path = excel_path
        self.excel_path_local = excel_path
        self.excel_read = [""]
        self.current_wb = None
        self.current_ws = None
        self.apply_all_policy = None # 덮어쓰기 정책을 위한 멤버 변수

        self.setWindowTitle("다중 케이스 실행 설정")
        self.resize(1000, 500)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 10, 10, 10)
        layout.setSpacing(8)

        lbl_ref = QLabel("참조파일: " + ref_name)
        lbl_ref.setWordWrap(True)
        lbl_ref.setStyleSheet("font-weight: bold; color: #1F2937; font-size: 13px;")
        layout.addWidget(lbl_ref)

        self.btn_excel = QPushButton("LoadCase 엑셀: " + os.path.basename(excel_path))
        self.btn_excel.setMinimumHeight(28)
        self.btn_excel.setStyleSheet("text-align: left; padding-left: 8px;")
        self.btn_excel.clicked.connect(self.on_excel_button_clicked)
        layout.addWidget(self.btn_excel)

        self.combo_sheet = QComboBox()
        self.combo_sheet.setMinimumHeight(28)
        self.combo_sheet.currentTextChanged.connect(self.load_sheet)
        layout.addWidget(self.combo_sheet)

        self.table = QTableWidget()
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionMode(QAbstractItemView.MultiSelection)
        self.table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.table, stretch=1)

        self.btn_reload = QPushButton("Excel reload")
        self.btn_reload.setMinimumSize(160, 34)
        self.btn_reload.setStyleSheet("font-weight: bold; background-color: #64748B; color: white; border-radius: 4px;")
        self.btn_reload.clicked.connect(lambda: self.reload_excel(self.excel_path_local, keep_selection=True))
        layout.addWidget(self.btn_reload)

        self.btn_run = QPushButton("Make file multi case run")
        self.btn_run.setMinimumSize(160, 34)
        self.btn_run.setStyleSheet("font-weight: bold; background-color: #1E40AF; color: white; border-radius: 4px;")
        self.btn_run.clicked.connect(self.on_make_run)
        layout.addWidget(self.btn_run)

        self.reload_excel(self.excel_path, keep_selection=False)

    @staticmethod
    def show_window(parent, text):
        """ 우클릭 참조 텍스트(text)에서 파일경로를 추출해 모달리스 창을 띄웁니다. """
        file_path = text.split(":", 1)[1].strip() if ":" in text else text.strip()
        if not os.path.exists(file_path):
            QMessageBox.warning(parent, "입력 파일 오류",
                                 "실행할 파일이 경로에 존재하지 않습니다.\n\n경로: " + file_path)
            return

        ref_dir = os.path.dirname(file_path)
        ref_name = os.path.basename(file_path)

        excel_path = ""
        for d in (ref_dir, os.path.dirname(ref_dir)):
            if not d or not os.path.isdir(d):
                continue
            for fn in os.listdir(d):
                if fn.lower().startswith("designloadcasetable") and fn.lower().endswith((".xlsx", ".xlsm", ".xls")):
                    excel_path = os.path.join(d, fn)
                    break
            if excel_path:
                break

        if not excel_path or not os.path.exists(excel_path):
            QMessageBox.critical(
                parent, "LoadCase 엑셀 없음",
                "참조파일 경로 또는 상위 디렉토리에서\n'DesignLoadCaseTable' 엑셀 파일을 찾을 수 없습니다.\n\n참조파일: " + ref_name
            )
            return

        print(f"[MultiCase] 엑셀 파일 활용: {excel_path}")
        # 클래스 자신을 인스턴스화하여 보여줍니다.
        MultiCaseRunWindow(parent, file_path, ref_name, excel_path).show()

    def _make_temp_copy(self, src_path):
        try:
            tmpf = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
            shutil.copyfile(src_path, tmpf.name)
            tmpf.close()
            return tmpf.name
        except Exception as e:
            QMessageBox.critical(self, "엑셀 복사 실패", str(e))
            return ""

    CHECKED_BG = "#D1FAE5"

    def apply_col_color(self, c, checked):
        bg = QBrush(QColor(self.CHECKED_BG)) if checked else QBrush()
        for r in range(self.table.rowCount()):
            item = self.table.item(r, c)
            if item is not None:
                item.setBackground(bg)

    @staticmethod
    def _col_letter_to_index(letter):
        letter = letter.upper()
        result = 0
        for char in letter:
            result = result * 26 + (ord(char) - ord('A') + 1)
        return result - 1

    def load_sheet(self, sheet_name):
        if not sheet_name or not self.excel_read[0]:
            return
        print(f"[MultiCase] load_sheet 호출: {sheet_name} / excel_read={self.excel_read[0]}")
        try:
            wb = openpyxl.load_workbook(self.excel_read[0], data_only=True)
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))

            self.current_wb = wb
            self.current_ws = ws

            hidden_rows = set()
            for r in ws.row_dimensions:
                if ws.row_dimensions[r].hidden:
                    hidden_rows.add(r)

            hidden_cols = set()
            for c in ws.column_dimensions:
                if ws.column_dimensions[c].hidden:
                    hidden_cols.add(c)

            wb.close()
        except Exception as e:
            print(f"[MultiCase] 시트 읽기 예외: {e}")
            rows = [["엑셀 읽기 실패:", str(e)]]
            hidden_rows = set()
            hidden_cols = set()
            ws = None
        if not rows:
            rows = [["(빈 시트)"]]

        data_cols = max((len(r) for r in rows if r), default=0)
        sheet_cols = ws.max_column if ws is not None else data_cols
        n_cols = max(data_cols, sheet_cols)

        self.table.setColumnCount(n_cols)
        self.table.setRowCount(len(rows))

        for r in range(len(rows)):
            self.table.showRow(r)
        for c in range(n_cols):
            self.table.showColumn(c)

        for r, row in enumerate(rows):
            for c in range(n_cols):
                val = row[c] if c < len(row) and row[c] is not None else ""
                self.table.setItem(r, c, QTableWidgetItem(str(val)))

        for excel_row in hidden_rows:
            table_row = excel_row - 1
            if 0 <= table_row < len(rows):
                self.table.hideRow(table_row)

        for col_letter in hidden_cols:
            col_idx = self._col_letter_to_index(col_letter)
            if 0 <= col_idx < n_cols:
                self.table.hideColumn(col_idx)

        self.table.resizeRowsToContents()
        print(f"[MultiCase] 테이블 세팅 완료: 데이터 {len(rows)}행")

    def __on_item_changed(self, item):
        r = item.row()
        checked = item.checkState() == Qt.Checked
        #print(f"[MultiCase] 행 {r} 체크 상태: {'ON' if checked else 'OFF'}")

    def reload_excel(self, src_path, keep_selection=False):
        if not src_path or not os.path.exists(src_path):
            QMessageBox.warning(self, "엑셀 없음", "선택한 엑셀 파일을 찾을 수 없습니다.")
            return
        prev_sheet = self.combo_sheet.currentText() if keep_selection else ""
        tmp = self._make_temp_copy(src_path)
        if not tmp:
            return
        self.excel_read[0] = tmp
        self.excel_path_local = src_path
        self.btn_excel.setText("LoadCase 엑셀: " + os.path.basename(src_path))
        try:
            wb = openpyxl.load_workbook(tmp)
            sheet_names = wb.sheetnames
            wb.close()
            self.combo_sheet.blockSignals(True)
            self.combo_sheet.clear()
            self.combo_sheet.addItems(sheet_names)
            if keep_selection and prev_sheet in sheet_names:
                self.combo_sheet.setCurrentText(prev_sheet)
                target = prev_sheet
            elif sheet_names:
                target = sheet_names[0]
            else:
                target = ""
            self.combo_sheet.blockSignals(False)
            if target:
                self.load_sheet(target)
        except Exception as e:
            print(f"[MultiCase] 시트목록 읽기 예외: {e}")
            QMessageBox.critical(self, "엑셀 읽기 실패", str(e))

    def on_excel_button_clicked(self):
        start_dir = os.path.dirname(self.excel_path_local) if self.excel_path_local else ""
        picked, _ = QFileDialog.getOpenFileName(
            self, "LoadCase 엑셀 선택", start_dir,
            "Excel Files (*.xlsx *.xlsm *.xls);;All Files (*)"
        )
        if picked:
            self.reload_excel(picked)

    def on_make_run(self):
        selected_rows = sorted({item.row() for item in self.table.selectedItems()})

        if self.table.rowCount() <= 1:
            print("[MultiCase] 처리할 데이터 행이 없습니다.")
            return

        if not selected_rows:
            print("[MultiCase] 선택된 행 없음 -> 화면에 보이는 행 내용을 기준으로 처리합니다.")
            selected_rows = [
                r for r in range(1, self.table.rowCount())
                if not self.table.isRowHidden(r)
                and any(
                    self.table.item(r, c) is not None and self.table.item(r, c).text().strip()
                    for c in range(self.table.columnCount())
                )
            ]

        headers = []
        for c in range(self.table.columnCount()):
            item = self.table.item(0, c)
            header_text = item.text().strip() if item is not None and item.text().strip() else f"Col_{c}"
            headers.append(header_text)

        display_to_param_key = {
            info.get("Description", ""): key
            for key, info in self.parameter_map.items()
            if info.get("Description")
        }
        print(f" display_to_param_key: {display_to_param_key}")

        changed_map = []
        of_run_list = {}

        for r in selected_rows:
            if r <= 0:
                continue

            row_values = []
            for c in range(self.table.columnCount()):
                item = self.table.item(r, c)
                value = item.text() if item is not None else ""
                row_values.append(value)

            row_data = {}
            for c, header in enumerate(headers):
                row_data[header] = row_values[c] if c < len(row_values) else ""

            of_word = str(row_data.get("OF_Word", "")).strip()
            file_symbol = str(row_data.get("File_Symbol", "")).strip()

            if of_word and file_symbol and of_word in self.parameter_map:
                old_symbol = self.parameter_map[of_word].get("File_Symbol", "")
                if old_symbol != file_symbol:
                    reply = QMessageBox.question(
                        self,
                        "parameter_map 변경 확인",
                        f"'{of_word}'의 File_Symbol을\n'{old_symbol}' -> '{file_symbol}'로 변경할까요?",
                        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                        QMessageBox.StandardButton.No,
                    )
                    if reply == QMessageBox.StandardButton.Yes:
                        self.parameter_map[of_word]["File_Symbol"] = file_symbol
                        changed_map.append(f"{of_word}: {old_symbol} -> {file_symbol}")
                    else:
                        print(f"[MultiCase] 변경 취소: {of_word} ({old_symbol} -> {file_symbol})")

            row_output = {}
            for header, value in row_data.items():
                if not value:
                    continue
                if header in {"Description", "OF_Word", "File_Symbol"}:
                    continue

                param_key = display_to_param_key.get(header)
                if not param_key:
                    continue

                row_output[param_key] = value

            first_value = str(row_values[0]).strip() if row_values else ""
            if first_value.isdigit():
                of_run_list[f"N{first_value}"] = row_output

        if changed_map:
            print("[MultiCase] parameter_map 변경사항:")
            for msg in changed_map:
                print(f"  - {msg}")

        print("\n" + "=" * 60 + "OF_RunList JSON Output:")
        print(json.dumps(of_run_list, ensure_ascii=False, indent=2))

        print("\n" + "=" * 60 + "Generated Filenames:")
        # 가장 큰 숫자의 자릿수 파악
        max_num = max((int(k[1:]) for k in of_run_list.keys()), default=0)
        num_width = len(str(max_num))

        # 각 파라미터별 최대 값 길이 파악
        param_max_length = {}
        for run_key, row_output in of_run_list.items():
            for param_key, value in row_output.items():
                if param_key in self.parameter_map:
                    value_str = str(value).strip()
                    if param_key not in param_max_length:
                        param_max_length[param_key] = 0
                    param_max_length[param_key] = max(param_max_length[param_key], len(value_str))
            
        for run_key, row_output in of_run_list.items():
            # run_key = "N1", row_output = {"WaveHs": "5.0", "WaveTp": "10.0", ...}
            num_part = run_key[1:]  # "1" 추출
            padded_num = num_part.zfill(num_width)  # "001" 등으로 padding
            filename_parts = [f"N{padded_num}"]

            for param_key in row_output.keys():
                if param_key in self.parameter_map:
                    file_symbol = self.parameter_map[param_key].get("File_Symbol", "")
                    truncated_symbol = file_symbol[:5] if file_symbol else ""
                    value = str(row_output[param_key]).strip().replace(" ", "")
                    if not value:
                        continue

                    # 숫자면 zero-padding, 아니면 그대로
                    numeric_check = value.replace('.', '').replace('-', '')
                    if numeric_check.isdigit():
                        padded_value = value.zfill(param_max_length.get(param_key, 0) ) if param_max_length.get(param_key, 0)  > 0 else value
                    else:
                        padded_value = value

                    if truncated_symbol:
                        filename_parts.append(f"{truncated_symbol}{padded_value}")

            filename = "_".join(filename_parts)
            of_run_list[run_key]["FST_File"] = filename
            print(f"  {filename}")

        # === 파일 생성 단계 (OpenFastIO.current_config 분석 및 중복 없이 카피) ===
        print("\n" + "=" * 60 + "파일 생성 및 저장 시작")

        main_fst_path = self.file_path
        main_fst_dir = os.path.dirname(main_fst_path)
        main_fst_basename = os.path.basename(main_fst_path)
        main_fst_base = os.path.splitext(main_fst_basename)[0]

        # 1. 복사 대상 하위 파일 수집 (OpenFastIO.current_config 활용)
        from src.core.openfast_io import OpenFastIO
        files_to_copy = []  # (key, 원래절대경로) 저장

        # 소프트 카피 시 고유 확장자(.ela, .aer 등)를 부여하고 물리 이사할 핵심 1차 파일 목록 정의
        soft_copy_targets = [
            "EDFile", "AeroFile", "ServoFile", "SeaStFile", "HydroFile", "MooringFile", "SubFile", "IceFile", "SoilFile", "InflowFile"
        ]        

        for key, info in OpenFastIO.current_config.items():
            if key not in soft_copy_targets or not isinstance(info, dict):
                continue
            
            raw_value = info.get("current") or info.get("default")
            if not raw_value:
                continue

            paths_to_check = []
            if isinstance(raw_value, list):
                paths_to_check.extend(raw_value)
            elif isinstance(raw_value, str):
                paths_to_check.append(raw_value)

            for path_item in paths_to_check:
                if not isinstance(path_item, str) or not path_item.strip():
                    continue
                # 숫자 파라미터 제외 가드코드
                if path_item.strip() in ["0", "1", "2", "3"] or path_item.replace('.', '', 1).isdigit():
                    continue
                if '.' not in os.path.basename(path_item):
                    continue

                if not os.path.isabs(path_item):
                    # 메인 .fst 파일이 위치한 원본 디렉토리를 기준으로 절대 경로화
                    resolved_path = os.path.normpath(os.path.join(os.path.dirname(main_fst_path), path_item))
                else:
                    resolved_path = os.path.normpath(path_item)

                if os.path.exists(resolved_path) and os.path.isfile(resolved_path):
                    if (key, resolved_path) not in files_to_copy:
                        files_to_copy.append((key, resolved_path))

        # config 템플릿의 정확한 Key 풀네임을 기준으로 고정 확장자 매핑 정의
        ext_rules = {
            "edfile": ".ela",
            "aerofile": ".aer",
            "servofile": ".ser",
            "seastfile": ".sea",
            "hydrofile": ".hyd",
            "mooringfile": ".moo",
            "subfile": ".sub",
            "icefile": ".ice",
            "soilfile": ".sol",
            "inflowfile": ".inf"
        }

        self.apply_all_policy = None # 덮어쓰기 정책 초기화

        # 2. 각 run별 생성 + 매핑
        for run_key, run_data in of_run_list.items():
            fst_filename = run_data.get("FST_File", "")
            if not fst_filename:
                continue
            
            main_replace_map = {}
            sub_replace_map = {}
            copied_files = []
            skip_this_run = False

            # 2-1. 메인 .fst 복사
            new_fst_path = os.path.join(main_fst_dir, f"{fst_filename}.fst")
            confirm = self.confirm_overwrite_with_policy(new_fst_path)
            if confirm is None:
                print("[MultiCase] 사용자가 작업을 취소했습니다.")
                return
            if not confirm:
                print(f"  [건너뜀] {fst_filename}.fst 이미 존재")
                skip_this_run = True

            if not skip_this_run:
                try:
                    shutil.copy2(main_fst_path, new_fst_path)
                    print(f"  [생성] {fst_filename}.fst")
                    of_run_list[run_key]["NewFST_Path"] = new_fst_path
                except Exception as e:
                    print(f"  [오류] {fst_filename}.fst 생성 실패: {e}")
                    continue

                # 2-2. 하위 파일 복사 및 매핑 (스마트 분석 수집 대상 기준)
                for key, src_path in files_to_copy:
                    src_basename = os.path.basename(src_path)
                    src_base = os.path.splitext(src_basename)[0]
                    src_ext = os.path.splitext(src_basename)[1]
                    
                    target_ext = ext_rules.get(key.lower(), src_ext)
                    new_sub_basename = f"{fst_filename}{target_ext}"
                    dst_path = os.path.join(main_fst_dir, new_sub_basename)

                    confirm = self.confirm_overwrite_with_policy(dst_path)
                    if confirm is None:
                        print("[MultiCase] 사용자가 작업을 취소했습니다.")
                        return
                    if not confirm:
                        print(f"  [건너뜀] {new_sub_basename} 이미 존재")
                        continue

                    try:
                        shutil.copy2(src_path, dst_path)
                        print(f"    → {new_sub_basename}")
                        copied_files.append(dst_path)

                        main_replace_map[src_basename] = new_sub_basename
                        main_replace_map[src_base] = fst_filename
                        sub_replace_map[src_basename] = new_sub_basename
                    except Exception as e:
                        print(f"    ✗ {new_sub_basename} 생성 실패: {e}")

                # 2-3. 내부 경로 참조 업데이트
                target_files = [new_fst_path] + copied_files
                sorted_main_names = sorted(main_replace_map.keys(), key=len, reverse=True)
                sorted_sub_names = sorted(sub_replace_map.keys(), key=len, reverse=True)

                for file_path in target_files:
                    if not os.path.exists(file_path):
                        continue

                    is_main_fst = os.path.abspath(file_path) == os.path.abspath(new_fst_path)
                    try:
                        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                            content = f.read()

                        modified = False
                        if is_main_fst:
                            loop_names = sorted_main_names
                        else:
                            current_base = os.path.splitext(os.path.basename(file_path))[0]
                            loop_names = [
                                name for name in sorted_sub_names
                                if not name.startswith(current_base)
                            ]

                        for old_name in loop_names:
                            if not old_name.strip():
                                continue
                            if old_name not in content:
                                continue

                            new_name = (
                                main_replace_map[old_name]
                                if is_main_fst
                                else sub_replace_map[old_name]
                            )
                            content = content.replace(old_name, new_name)
                            modified = True

                        if modified:
                            with open(file_path, 'w', encoding='utf-8') as f:
                                f.write(content)
                    except Exception as e:
                        print(f"    ✗ {os.path.basename(file_path)} 내용 업데이트 실패: {e}")

                # 2-4. .fst 하위 파일들 KEY Update (변수들을 Multi에서 생성된 값으로 지정)
                print(f"\n🚀 [Processing] {run_key} 세트 내부 KEY 데이터 변환 작업 시작")
                print(f"  → 대상 파일명: {fst_filename}.fst")

                new_data = {}
                new_data_specified = {}

                # run_data 내부에 들어있는 모든 파라미터들(WaveHs, WaveTp, NacYaw 등)을 순회하며 매핑
                for param_key, param_value in run_data.items():
                    # 파일명 제어용 내부 메타 데이터 키들은 치환 대상에서 제외합니다.
                    if param_key in ["FST_File", "NewFST_Path"]:
                        continue
                    
                    clean_val = str(param_value).strip()
                    if not clean_val:
                        continue

                    if param_key in self.parameter_map:
                        target_ext = self.parameter_map[param_key].get("Target_File", "").lower()
                        if not target_ext:
                            continue

                        # [부호 정제 규칙 적용] 맨 앞의 '+' 기호는 제거하고 '-' 기호는 보존합니다. 예시: "+8" -> "8", "-8" -> "-8"
                        if clean_val.startswith("+"):
                            clean_val = clean_val[1:]  # 맨 앞 한 글자(+)를 제외한 나머지 문자열 추출
                        elif clean_val.startswith("-"):
                            pass  # 마이너스 기호는 조건에 맞춰 그대로 유지합니다.

                        # 값에 부호나 텍스트 문자가 섞여있다면 쌍따옴표 래핑, 순수 숫자면 그대로 주입
                        if clean_val.replace('.', '', 1).replace('-', '', 1).isdigit():
                            formatted_val = clean_val
                        else:
                            formatted_val = f'"{clean_val}"'

                        # 해당 확장자 묶음 데이터셋에 KEY-VALUE 누적 저장
                        if target_ext not in new_data_specified:
                            new_data_specified[target_ext] = {}
                        new_data_specified[target_ext][param_key] = formatted_val


                if ".fst" not in new_data_specified:
                    new_data_specified[".fst"] = {}

                for key, src_path in files_to_copy:
                    src_basename = os.path.basename(src_path)
                    new_sub_basename = main_replace_map.get(src_basename)
                    if new_sub_basename:
                        new_data_specified[".fst"][key] = f'"{new_sub_basename}"'

                print(f"  → new_data_specified = {new_data_specified}")

                # 최종 실행 단계: 각 확장자별로 표적 파일을 찾아 매칭되는 new_data만 주입 실행
                for target_ext, new_data in new_data_specified.items():
                    if not new_data:
                        continue

                    # 현재 수집된 target_files 중 해당 확장자를 가진 실제 물리 경로를 색출합니다.
                    actual_file_path = os.path.join(main_fst_dir, fst_filename + target_ext)
                    
                    print(f"  fst_filename = {fst_filename} \n  target_ext = {target_ext}\n  actual_file_path={actual_file_path}\n")

                    if target_ext == ".inf" and "FileName_BTS" in new_data_specified[".inf"]:
                            # 기존에 담겨있던 원본 값(예: '"s6"' 또는 "s6") 추출
                            original_bts = new_data_specified[".inf"]["FileName_BTS"]
                            clean_bts    = original_bts.replace('"', '').strip()

                            if not self.current_ws: 
                                print("❌ [오류] 참조할 활성화된 엑셀 시트(self.current_ws)가 없습니다.")
                                return
                            
                            # 확장자 결합 및 큰따옴표 포맷팅 (예: "s6" + ".bts" -> '"s6.bts"')
                            new_data_specified[".inf"]["FileName_BTS"] = f'"../../3_Wind/DLC{self.current_ws.title}/{clean_bts}.bts"'
                    
                    if actual_file_path and os.path.exists(actual_file_path):
                        try:                          
                            OpenFastIO.save_module_data(actual_file_path, new_data)
                            
                            print(f"------ {target_ext} 파일 내부 변수 치환 동기화 완료\n")
                        except Exception as e:
                            print(f"✗ {target_ext} 파일 대상 데이터 주입 중 에러 발생: {e}")
                    else:
                        # .inf 파일 등이 누락되었을 때의 경고 안내 메시지
                        print(f"⚠️ [경로 누락] {target_ext} 타입에 해당하는 생성된 물리 파일 경로를 찾을 수 없어 치환을 스킵합니다.")

                # 대량의 run_key 루프 연산 중 PyQt UI 화면이 먹통(Freeze)이 되지 않도록 스레드 이벤트 비우기
                QCoreApplication.processEvents()

        print("=" * 60 + "\n")

    def confirm_overwrite_with_policy(self, path):
        """ helper: 기존 파일 존재시 사용자에게 덮어쓰기 확인 """
        if not os.path.exists(path):
            return True  # 새 파일이면 무조건 생성
            
        # 이미 "모두 적용" 정책이 결정되어 있다면 팝업 없이 즉시 반환
        if self.apply_all_policy == "YES_ALL":
            return True
        if self.apply_all_policy == "NO_ALL":
            return False

        # Custom 메시지 박스 생성 (기본 메시지 박스는 YesToAll 버튼 커스텀이 까다로움)
        msg_box = QMessageBox(self)
        msg_box.setWindowTitle("파일 덮어쓰기 확인")
        msg_box.setText(f"이미 존재하는 파일입니다:\n{os.path.basename(path)}\n\n덮어쓰시겠습니까?")
        
        # 필요한 버튼들 등록
        yes_btn = msg_box.addButton("예 (&Y)", QMessageBox.ButtonRole.YesRole)
        yes_all_btn = msg_box.addButton("모두 예 (&A)", QMessageBox.ButtonRole.YesRole)
        no_btn = msg_box.addButton("아니오 (&N)", QMessageBox.ButtonRole.NoRole)
        no_all_btn = msg_box.addButton("모두 아니오 (&L)", QMessageBox.ButtonRole.NoRole)
        cancel_btn = msg_box.addButton(QMessageBox.StandardButton.Cancel)
        
        msg_box.setDefaultButton(no_btn) # 엔터 시 안전하게 '아니오'가 선택되도록 설정
        msg_box.exec_()
        
        clicked_btn = msg_box.clickedButton()
        
        if clicked_btn == yes_btn:
            return True
        elif clicked_btn == yes_all_btn:
            self.apply_all_policy = "YES_ALL" # 💡 이후 파일들은 팝업 없이 모두 True
            return True
        elif clicked_btn == no_btn:
            return False
        elif clicked_btn == no_all_btn:
            self.apply_all_policy = "NO_ALL"  # 💡 이후 파일들은 팝업 없이 모두 False
            return False
        else:
            return None # Cancel (전체 취소)



    def closeEvent(self, event):
        if self.excel_read[0] and os.path.exists(self.excel_read[0]):
            try:
                os.remove(self.excel_read[0])
            except Exception:
                pass
        super().closeEvent(event)
